"""Read-only CSV/XLSX input; conservative validation and auditable reporting."""
import argparse, csv, json, re, hashlib
from pathlib import Path
from datetime import datetime, date
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from collections import defaultdict, Counter
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT=Path(__file__).resolve().parents[1]
CONFIG=json.loads((ROOT/'src/field-mapping.json').read_text())
ORDER_FIELDS='channel order_id line_id order_date country sku quantity unit_price_usd discount_rate customer_type traffic_source'.split()
PRODUCT_FIELDS='sku product_name category'.split()
REFUND_FIELDS='refund_id channel order_id line_id refund_date refund_amount_usd'.split()

def money(value):
    s=str(value).strip().replace('$','').replace(',','')
    if not s:raise ValueError('missing_amount')
    v=Decimal(s)
    if not v.is_finite() or v<0:raise ValueError('invalid_amount')
    return v

def cents(v):return int((v*100).quantize(Decimal('1'),rounding=ROUND_HALF_UP))

def parse_date(value):
    s=str(value).strip()
    if not re.fullmatch(r'\d{4}[-/]\d{2}[-/]\d{2}',s):raise ValueError('ambiguous_or_invalid_date')
    return date.fromisoformat(s.replace('/','-')).isoformat()

def dump_csv(path,rows,fields):
    path.parent.mkdir(parents=True,exist_ok=True)
    with path.open('w',encoding='utf-8',newline='') as f:
        w=csv.DictWriter(f,fieldnames=fields,extrasaction='ignore');w.writeheader();w.writerows(rows)

def read_input(path, fields, kind, ledger, fixes):
    records=[]
    for file in sorted(path):
        if file.suffix=='.csv':
            with file.open(encoding='utf-8-sig',newline='') as f:
                reader=csv.DictReader(f); headers=reader.fieldnames or []; raw=list(reader)
        else:
            df=pd.read_excel(file,dtype=str,keep_default_na=False);headers=list(df.columns);raw=df.to_dict('records')
        canonical=[CONFIG['aliases'].get(k,k) for k in headers]
        missing=set(fields)-set(canonical)
        if missing or len(set(canonical))!=len(canonical):raise ValueError(f'{file.name}: missing fields {sorted(missing)} or duplicate mapped headers; expected {fields}')
        for line,r in enumerate(raw,2):
            row={CONFIG['aliases'].get(k,k):str(v).strip() for k,v in r.items()}
            record={'kind':kind,'file':file.name,'row':line,'original':r,'value':row,'status':'retained','reason':'valid'}
            ledger.append(record);records.append(record)
            for k,v in r.items():
                key=CONFIG['aliases'].get(k,k)
                if k!=key or str(v)!=row[key]:fixes.append(dict(file=file.name,row=line,field=key,before=str(v),after=row[key],rule='N01-field-alias-or-whitespace'))
    return records

def flag(record,reason):record['status']='pending';record['reason']=reason

def dedupe(records,keys,fields):
    groups=defaultdict(list)
    for r in records:
        if all(r['value'].get(k,'') for k in keys):groups[tuple(r['value'].get(k,'') for k in keys)].append(r)
    for group in groups.values():
        variants={tuple(r['value'].get(k,'') for k in fields) for r in group}
        if len(variants)>1:
            for r in group:flag(r,'K02-same-key-conflict')
        else:
            if all(r['status']=='retained' for r in group):
                for r in group[1:]:r['status']='deduplicated';r['reason']='K01-exact-business-duplicate'

def normalize(record,fixes):
    r=record['value']
    try:
        for k in list(r):
            before=r[k]
            if k=='channel':r[k]=CONFIG['channels'][before.lower()]
            elif k=='country':
                r[k]=before.upper()
                if r[k] not in CONFIG['countries']:raise ValueError('invalid_country')
            elif k=='sku':r[k]=before.upper()
            elif k=='category':r[k]=CONFIG['categories'][before.lower()]
            elif k=='customer_type':r[k]=CONFIG['customers'][before.lower()]
            elif k=='traffic_source':r[k]=CONFIG['traffic'][before.lower()]
            elif k in ['order_date','refund_date']:r[k]=parse_date(before)
            elif k in ['unit_price_usd','refund_amount_usd']:
                v=money(before)
                if v*100!=(v*100).to_integral_value():raise ValueError('amount_has_subcent_precision')
                r[k]=str(v.quantize(Decimal('.01')))
            elif k=='quantity':
                if not re.fullmatch(r'[1-9]\d*',before):raise ValueError('invalid_quantity')
                r[k]=str(int(before))
            elif k=='discount_rate':
                v=Decimal(before)
                if not v.is_finite() or not 0<=v<=1:raise ValueError('invalid_discount')
                r[k]=str(v.normalize())
            elif not before:raise ValueError('missing_'+k)
            if r[k]!=before:fixes.append(dict(file=record['file'],row=record['row'],field=k,before=before,after=r[k],rule='N02-approved-normalization'))
    except (ValueError,InvalidOperation,KeyError) as e:flag(record,'V01-'+str(e))

def metrics(rows,refunds):
    keys={(r['channel'],r['order_id'],r['line_id']) for r in rows};orders={(r['channel'],r['order_id']) for r in rows}
    rr=[r for r in refunds if (r['channel'],r['order_id'],r['line_id']) in keys]
    sales=sum(r['sales_cents'] for r in rows);ref=sum(r['refund_cents'] for r in rr)
    positive={(r['channel'],r['order_id']) for r in rr if r['refund_cents']>0}
    returning={(r['channel'],r['order_id']) for r in rows if r['customer_type']=='Returning'}
    return dict(sales_cents=sales,refund_cents=ref,net_cents=sales-ref,orders=len(orders),lines=len(rows),refunded_orders=len(positive),returning_orders=len(returning),aov_cents=sales/len(orders) if orders else None,refund_rate=len(positive)/len(orders) if orders else None,returning_share=len(returning)/len(orders) if orders else None)

def build(input_path,output):
    ledger=[];fixes=[];files=list(input_path.glob('*.csv'))+list(input_path.glob('*.xlsx'))
    products=read_input([f for f in files if f.stem=='products'],PRODUCT_FIELDS,'products',ledger,fixes)
    refunds=read_input([f for f in files if f.stem=='refunds'],REFUND_FIELDS,'refunds',ledger,fixes)
    orders=read_input([f for f in files if f.stem not in ['products','refunds']],ORDER_FIELDS,'orders',ledger,fixes)
    if not orders or not products or not refunds:raise ValueError('Input requires order exports, products and refunds tables')
    for r in ledger:normalize(r,fixes)
    dedupe(products,['sku'],PRODUCT_FIELDS);dedupe(orders,['channel','order_id','line_id'],ORDER_FIELDS);dedupe(refunds,['channel','refund_id'],REFUND_FIELDS)
    product_map={r['value']['sku']:r['value'] for r in products if r['status']=='retained'}
    for r in orders:
        if r['status']=='retained' and r['value']['sku'] not in product_map:flag(r,'J01-unknown-or-conflicted-SKU')
    order_groups=defaultdict(list)
    for r in orders:order_groups[(r['value'].get('channel',''),r['value'].get('order_id',''))].append(r)
    for group in order_groups.values():
        active=[r for r in group if r['status']=='retained']
        attrs={tuple(r['value'].get(k,'') for k in ['order_date','country','customer_type','traffic_source']) for r in active}
        if len(attrs)>1:
            for r in active:flag(r,'O01-order-attribute-conflict')
        if any(r['status']=='pending' for r in group):
            for r in group:
                if r['status']=='retained':flag(r,'O02-whole-order-excluded')
    clean=[]
    for record in orders:
        if record['status']!='retained':continue
        r=dict(record['value']);r.update(product_map[r['sku']]);r['quantity']=int(r['quantity']);r['sales_cents']=cents(Decimal(r['quantity'])*money(r['unit_price_usd'])*(1-Decimal(r['discount_rate'])));r.update(source_file=record['file'],source_row=record['row']);clean.append(r)
    line_map={(r['channel'],r['order_id'],r['line_id']):r for r in clean}
    refund_groups=defaultdict(list)
    for record in refunds:
        if record['status']!='retained':continue
        r=record['value'];key=(r['channel'],r['order_id'],r['line_id'])
        if key not in line_map:flag(record,'R01-unmatched-valid-sale');continue
        if r['refund_date']<line_map[key]['order_date']:flag(record,'R02-refund-before-sale');continue
        refund_groups[key].append(record)
    # Over-refund makes the whole linked refund group pending; never choose an arbitrary refund.
    for key,group in refund_groups.items():
        if sum(cents(money(r['value']['refund_amount_usd'])) for r in group)>line_map[key]['sales_cents']:
            for r in group:flag(r,'R03-cumulative-refund-exceeds-sale')
    clean_ref=[]
    for record in refunds:
        if record['status']=='retained':
            r=dict(record['value']);r['refund_cents']=cents(money(r['refund_amount_usd']));r.update(source_file=record['file'],source_row=record['row']);clean_ref.append(r)
    clean_products=[r['value'] for r in products if r['status']=='retained']
    quality={kind:dict(Counter(r['status'] for r in ledger if r['kind']==kind)) for kind in ['orders','products','refunds']}
    quality['excluded_orders']=len({(r['value'].get('channel'),r['value'].get('order_id')) for r in orders if r['status']=='pending'})
    monthly=[]
    for month in sorted({r['order_date'][:7] for r in clean}):monthly.append({'month':month,**metrics([r for r in clean if r['order_date'].startswith(month)],clean_ref)})
    result=dict(rows=clean,refunds=clean_ref,quality=quality,monthly=monthly,totals=metrics(clean,clean_ref),coverage=[min(r['value']['order_date'] for r in orders if re.fullmatch(r'\d{4}-\d{2}-\d{2}',r['value']['order_date'])),max(r['value']['order_date'] for r in orders if re.fullmatch(r'\d{4}-\d{2}-\d{2}',r['value']['order_date']))],generated=datetime.now().astimezone().isoformat(),exceptions=[{'kind':r['kind'],'file':r['file'],'row':r['row'],'status':r['status'],'reason':r['reason'],'original':r['original']} for r in ledger if r['status']!='retained'])
    output.mkdir(parents=True,exist_ok=True);report=output/'report';report.mkdir(exist_ok=True)
    clean_path=ROOT/'data/clean' if output.resolve()==(ROOT/'deliverables').resolve() else output/'clean'
    dump_csv(clean_path/'orders.csv',clean,ORDER_FIELDS+['product_name','category','sales_cents','source_file','source_row']);dump_csv(clean_path/'products.csv',clean_products,PRODUCT_FIELDS);dump_csv(clean_path/'refunds.csv',clean_ref,REFUND_FIELDS+['refund_cents','source_file','source_row'])
    audit=[dict(kind=r['kind'],file=r['file'],row=r['row'],status=r['status'],reason=r['reason'],original=json.dumps(r['original'],ensure_ascii=False),normalized=json.dumps(r['value'],ensure_ascii=False)) for r in ledger]
    dump_csv(output/'row-ledger.csv',audit,list(audit[0]));dump_csv(output/'normalizations.csv',fixes,list(fixes[0]) if fixes else ['file','row','field','before','after','rule'])
    (report/'data.js').write_text('window.REPORT_DATA = '+json.dumps(result,ensure_ascii=False)+';',encoding='utf-8');(output/'metrics.json').write_text(json.dumps(result['totals'],indent=2),encoding='utf-8')
    wb=Workbook();wb.remove(wb.active)
    tabs={'Monthly summary':monthly,'Orders':clean,'Products':clean_products,'Refunds':clean_ref,'Exceptions':[a for a in audit if a['status']!='retained'],'Row ledger':audit,'Normalizations':fixes,'Definitions':[{'metric':'Sales / 销售额','definition':'USD merchandise only. quantity × unit price × (1 − discount), HALF_UP to cents per line.'},{'metric':'Refunds / 退款额','definition':'Valid linked refunds assigned to ORIGINAL SALE DATE, not cash-flow date.'},{'metric':'Net sales / 净销售额','definition':'Sales less refunds; excludes tax, shipping, fees and costs. Not profit.'},{'metric':'Orders / 订单数','definition':'Distinct channel + order ID; category counts are not additive.'},{'metric':'AOV / 平均订单金额','definition':'Selected sales / selected orders; category filters count selected merchandise only.'},{'metric':'Refund order rate / 退款订单率','definition':'Orders with positive linked refund / selected orders. Not item return rate.'},{'metric':'Returning share / 老客订单占比','definition':'Source-marked Returning orders / selected orders. Not retention.'},{'metric':'Project','definition':'Nestora Home fictional brand; synthetic Q3 2025 data.'}]}
    for idx,(name,rs) in enumerate(tabs.items()):
        ws=wb.create_sheet(name);fields=list(rs[0]) if rs else ['No records'];ws.append(fields)
        for r in rs:
            vals=[]
            for k in fields:
                v=r.get(k)
                if k in ['unit_price_usd','discount_rate','refund_amount_usd']:v=float(v) if v is not None else None
                elif k in ['order_date','refund_date']:v=datetime.fromisoformat(v) if v else None
                vals.append(v)
            ws.append(vals)
        ws.freeze_panes='A2';ws.sheet_view.showGridLines=False;ws.auto_filter.ref=ws.dimensions;ws.row_dimensions[1].height=32
        for cell in ws[1]:cell.fill=PatternFill('solid',fgColor='1E3A8A');cell.font=Font(color='FFFFFF',bold=True);cell.alignment=Alignment(wrap_text=True,vertical='center')
        for col,k in enumerate(fields,1):
            from openpyxl.utils import get_column_letter
            width=min(55,max(18,len(k)+3))
            if k in ['metric','reason','rule']:width=43
            if k=='definition':width=90
            if k in ['file','source_file']:width=30
            ws.column_dimensions[get_column_letter(col)].width=width
            for cells in ws.iter_rows(min_row=2,min_col=col,max_col=col):
                c=cells[0];c.alignment=Alignment(vertical='top',wrap_text=k in ['definition','reason','rule','metric']);c.font=Font(color='0F172A',size=11)
                if k.endswith('_cents'):c.number_format='#,##0'
                elif k in ['unit_price_usd','refund_amount_usd']:c.number_format='"$"#,##0.00'
                elif k in ['refund_rate','returning_share','discount_rate']:c.number_format='0.0%'
                elif k in ['order_date','refund_date']:c.number_format='yyyy-mm-dd'
        if rs:
            table=Table(displayName=f'Data{idx}',ref=ws.dimensions);table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);ws.add_table(table)
    wb.save(output/'cleaned-data.xlsx')
    counts='\n'.join(f'- {k}: {v}' for k,v in quality.items())
    reason_counts=dict(Counter(r['reason'] for r in ledger if r['status']=='pending'))
    (output/'quality-report.md').write_text(f'# Data quality\n\nSynthetic demo data. Unresolved rows are excluded.\n\n{counts}\n\nPending reasons: {json.dumps(reason_counts,indent=2)}\n\nEvery source row is accounted for in row-ledger.csv. Approved changes are in normalizations.csv. Refund overages exclude the complete refund group. Missing values are not imputed. See validation/ for independently generated controls. Human spot-check remains pending until performed by the owner.\n',encoding='utf-8')
    return result

if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--input',type=Path,required=True);ap.add_argument('--output',type=Path,required=True);args=ap.parse_args()
    result=build(args.input,args.output);print(json.dumps({'totals':result['totals'],'quality':result['quality']},indent=2))
